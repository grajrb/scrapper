"""CSV and Excel exports.

Excel workbook sheets (7): New Jobs, High Match, All Active Jobs,
Target Companies, Sources, Scrape Errors, ATS Discovery.
Columns match the agreed 15-column job layout.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

from .config import AppConfig
from .database import Database
from .models import Job, RemoteStatus, best_application_url
from .ranking import is_high_match, rank_jobs

COLUMNS = [
    "Company", "Priority", "Job Title", "Location", "Remote", "Experience",
    "Match Score", "Match Reasons", "Posted Date", "First Seen", "Source",
    "ATS", "Canonical URL", "Application URL", "Status",
]

SHEET_ORDER = [
    ("New Jobs", "new_jobs"),
    ("High Match", "high_match"),
    ("All Active Jobs", "all_active"),
    ("Target Companies", "companies"),
    ("Sources", "sources"),
    ("Scrape Errors", "errors"),
    ("ATS Discovery", "ats_discovery"),
]


@dataclass
class ExportContext:
    new_jobs: list[dict] = field(default_factory=list)
    high_match: list[dict] = field(default_factory=list)
    all_active: list[dict] = field(default_factory=list)
    companies: list[dict] = field(default_factory=list)
    sources: list[dict] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)
    ats_discovery: list[dict] = field(default_factory=list)


def _remote_display(job: Job) -> str:
    mapping = {
        RemoteStatus.REMOTE.value: "Yes (remote)",
        RemoteStatus.HYBRID.value: "Hybrid",
        RemoteStatus.ONSITE.value: "No (onsite)",
    }
    return mapping.get(job.remote_status, "Unknown")


def job_row(job: Job, priority: str) -> dict[str, Any]:
    """Map a Job to the export column layout."""
    reasons = list(job.match_reasons or [])
    if job.match_warnings:
        reasons += [f"WARNING: {w}" for w in job.match_warnings]
    return {
        "Company": job.company_name,
        "Priority": priority,
        "Job Title": job.title,
        "Location": job.location or "Not listed",
        "Remote": _remote_display(job),
        "Experience": job.experience_display,
        "Match Score": job.match_score if job.match_score is not None else "",
        "Match Reasons": "; ".join(reasons),
        "Posted Date": job.posted_date or "Not listed",
        "First Seen": job.first_seen,
        "Source": job.source,
        "ATS": job.ats or "",
        "Canonical URL": job.canonical_url or job.source_url or "",
        "Application URL": best_application_url(job) or "",
        "Status": job.status + (" [FIXTURE]" if job.is_fixture else ""),
    }


def build_export_context(db: Database, cfg: AppConfig,
                         error_limit: int = 200) -> ExportContext:
    companies = db.get_company_rows()
    priority_map = {c["company_name"]: c["priority"] for c in companies}
    jobs = db.list_jobs(active_only=True)
    ranked = rank_jobs(jobs, priority_map)

    all_rows = [job_row(j, priority_map.get(j.company_name, ""))
                for j in ranked]
    new_rows = [job_row(j, priority_map.get(j.company_name, ""))
                for j in ranked if j.status == "NEW"]
    high_rows = [job_row(j, priority_map.get(j.company_name, ""))
                 for j in ranked
                 if is_high_match(j, cfg.search.high_match_score)]

    return ExportContext(
        new_jobs=new_rows,
        high_match=high_rows,
        all_active=all_rows,
        companies=companies,
        sources=db.get_source_rows(),
        errors=db.get_errors(limit=error_limit),
        ats_discovery=db.get_ats_rows(),
    )


def export_csv(context: ExportContext, path: Path) -> Path:
    """Write ALL active jobs to one CSV (utf-8-sig for Excel compatibility)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS)
        writer.writeheader()
        for row in context.all_active:
            writer.writerow({k: row.get(k, "") for k in COLUMNS})
    return path


def _autosize(sheet, header_row: int = 1, max_width: int = 60) -> None:
    for column_cells in sheet.columns:
        letter = column_cells[0].column_letter
        width = max((len(str(c.value)) for c in column_cells
                     if c.value is not None), default=10)
        sheet.column_dimensions[letter].width = min(max(12, width + 2), max_width)
    sheet.freeze_panes = f"A{header_row + 1}"


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return value


def export_excel(context: ExportContext, path: Path) -> Path:
    """Write the 7-sheet workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, attr in SHEET_ORDER:
        rows: list[dict] = getattr(context, attr, []) or []
        ws = wb.create_sheet(title=sheet_name)
        if sheet_name in ("New Jobs", "High Match", "All Active Jobs"):
            headers = COLUMNS
        elif sheet_name == "Target Companies":
            headers = ["company_name", "company_domain", "careers_url",
                       "priority", "enabled", "ats"]
        elif sheet_name == "Sources":
            headers = ["source_key", "source_type", "enabled", "status",
                       "last_run_at", "last_success_at", "last_error",
                       "jobs_found"]
        elif sheet_name == "Scrape Errors":
            headers = ["occurred_at", "run_id", "source", "company", "url",
                       "error_type", "http_status", "message"]
        else:  # ATS Discovery
            headers = ["company_name", "careers_url", "ats", "confidence",
                       "status", "jobs_url", "discovered_at"]

        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([_cell(row.get(h)) for h in headers])
        _autosize(ws)

    wb.save(str(path))
    return path


def export_all(db: Database, cfg: AppConfig, fmt: str = "both",
               out_dir: Optional[Path] = None) -> list[Path]:
    """Export CSV and/or Excel into ``out_dir`` (default data/exports)."""
    context = build_export_context(db, cfg)
    target = Path(out_dir) if out_dir else cfg.export_dir
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    written: list[Path] = []

    if fmt in ("csv", "both"):
        written.append(export_csv(context, target / f"jobs_{stamp}.csv"))
    if fmt in ("xlsx", "excel", "both"):
        written.append(export_excel(context, target / f"jobs_{stamp}.xlsx"))
    return written

