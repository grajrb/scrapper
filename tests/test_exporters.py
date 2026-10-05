"""Tests for CSV and Excel exports."""
import csv

from openpyxl import load_workbook

from src.exporters import (
    COLUMNS,
    SHEET_ORDER,
    build_export_context,
    export_all,
    export_csv,
    export_excel,
    job_row,
)


def _seed(db, app_config, make_job):
    db.upsert_companies(app_config.companies)
    job = make_job(match_score=88, status="NEW",
                   match_reasons=["Exact role match: backend engineer"],
                   match_warnings=["Salary unavailable"])
    job.id = db.insert_job(job)
    db.add_job_source(job.id, job.source, job.source_job_id, job.source_url)
    db.ensure_sources({"fixture": "other"})
    db.log_error(1, source="indeed", error_type="MANUAL_REVIEW",
                 message="blocked")
    return job


def test_job_row_has_all_columns(make_job):
    row = job_row(make_job(match_score=88,
                           match_reasons=["Exact role match: backend engineer"],
                           match_warnings=["Salary unavailable"]), "A")
    assert set(row.keys()) == set(COLUMNS)
    assert row["Priority"] == "A"
    assert row["Match Score"] == 88
    assert row["Application URL"].startswith("https://")
    assert "WARNING: Salary unavailable" in row["Match Reasons"]


def test_csv_export(db, app_config, tmp_path, make_job):
    _seed(db, app_config, make_job)
    context = build_export_context(db, app_config)
    path = export_csv(context, tmp_path / "out.csv")
    with path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert rows, "expected at least one exported job"
    assert list(rows[0].keys()) == COLUMNS
    assert rows[0]["Company"] == "Acme Cloud"
    assert rows[0]["Match Score"] == "88"


def test_excel_export_seven_sheets(db, app_config, tmp_path, make_job):
    _seed(db, app_config, make_job)
    context = build_export_context(db, app_config)
    path = export_excel(context, tmp_path / "out.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == [name for name, _ in SHEET_ORDER]
    ws = wb["New Jobs"]
    assert ws.cell(row=1, column=1).value == "Company"
    assert ws.max_row >= 2
    assert wb["Target Companies"].max_row >= 2
    assert wb["Scrape Errors"].max_row >= 2


def test_export_all_writes_both_formats(db, app_config, tmp_path, make_job):
    _seed(db, app_config, make_job)
    written = export_all(db, app_config, fmt="both", out_dir=tmp_path)
    suffixes = {p.suffix for p in written}
    assert suffixes == {".csv", ".xlsx"}
    assert all(p.exists() for p in written)
